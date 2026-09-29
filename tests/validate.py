"""Offline validation of the M365 audit pipeline (no tenant needed).

What it does, end to end, using the real Excel installed on this PC:
  1. Runs Collect-M365Audit.ps1 -SampleData over reference/sample_data
     -> produces DATA\\*.csv + MANUAL\\*.csv + run.json.
  2. Copies dist\\M365_Audit_Dashboard.xlsm, types dummy license prices into
     Config (via Excel COM), then runs Refresh-Workbook.ps1 (the macro-free
     importer - same logic as the VBA module) against the copy.
  3. Reads back the recalculated values with openpyxl and asserts every key
     KPI against values computed independently with pandas.
  4. Asserts ZERO formula errors on every sheet.
  5. Repeats the import with header-only CSVs (empty tenant case): the
     dashboard must show zeros/dashes and still ZERO formula errors.
  6. Exports the Dashboard to PDF and asserts it is exactly one page.
  7. Static consistency checks: chart XML bound to D_ sheets; the dataset
     contract identical across contract.py / modImport.bas /
     Refresh-Workbook.ps1; SKU map identical between the collector and the
     workbook builder.

Run:  python tests\\validate.py     (from the dist folder)
Exit code 0 = all green.
"""
import csv
import datetime as dt
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
DIST = os.path.dirname(HERE)
ROOT = os.path.dirname(DIST)
sys.path.insert(0, os.path.join(DIST, "tools"))
import contract  # noqa: E402
from build_workbook import SKU_MAP  # noqa: E402

SAMPLE = os.path.join(ROOT, "reference", "sample_data")
XLSM = os.path.join(DIST, "M365_Audit_Dashboard.xlsm")
XLSX = os.path.join(DIST, "M365_Audit_Dashboard.xlsx")

FAILURES = []


def check(name, cond, detail=""):
    tag = "PASS" if cond else "FAIL"
    print(f"  [{tag}] {name}" + (f"  ({detail})" if detail and not cond else ""))
    if not cond:
        FAILURES.append(f"{name}: {detail}")


def run(cmd, **kw):
    print("  $", " ".join(cmd if isinstance(cmd, list) else [cmd]))
    r = subprocess.run(cmd, capture_output=True, text=True, **kw)
    if r.returncode != 0:
        print(r.stdout[-3000:])
        print(r.stderr[-2000:])
        raise RuntimeError(f"command failed ({r.returncode})")
    return r.stdout


def ps(script_path, *args):
    return run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass",
                "-File", script_path] + list(args))


# ---------------------------------------------------------------------------
# CSV loading that mirrors the VBA importer's typing (OpenText, Local:=False)
# ---------------------------------------------------------------------------
def load_csv(path, cols):
    rows = []
    with open(path, encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh)
        header = next(reader)
        assert [h.strip() for h in header] == [c for c, _t in cols], \
            f"{os.path.basename(path)} header mismatch: {header}"
        for raw in reader:
            row = {}
            for (name, typ), val in zip(cols, raw):
                if val == "":
                    row[name] = None
                elif typ == "B":
                    row[name] = val.upper() == "TRUE"
                elif typ == "N":
                    row[name] = float(val)
                elif typ == "D":
                    row[name] = dt.date.fromisoformat(val[:10])
                else:
                    row[name] = val
            rows.append(row)
    return pd.DataFrame(rows, columns=[c for c, _t in cols])


ERR_VALUES = {"#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NULL!", "#NUM!"}


def scan_errors(wb):
    errors = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value in ERR_VALUES:
                    errors.append(f"{ws.title}!{cell.coordinate}={cell.value}")
    return errors


def cell(wb, ref):
    sheet, coord = ref.split("!")
    return wb[sheet][coord].value


def num(v):
    if v is None:
        return 0.0
    return float(v)


def xlround(x, digits=0):
    """Excel ROUND: half away from zero (Python round() is banker's rounding)."""
    from decimal import Decimal, ROUND_HALF_UP
    q = Decimal(10) ** -digits
    return float(Decimal(str(x)).quantize(q, rounding=ROUND_HALF_UP))


# ---------------------------------------------------------------------------
# Step 1: sample collection
# ---------------------------------------------------------------------------
def step_collect(workdir):
    print("\n== 1. Collector -SampleData ==")
    ps(os.path.join(DIST, "Collect-M365Audit.ps1"),
       "-SampleData", SAMPLE, "-OutputPath", os.path.join(workdir, "DATA"))
    data = os.path.join(workdir, "DATA")
    for key, (fname, _s, _m, _c) in contract.DATASETS.items():
        check(f"{fname} produced", os.path.exists(os.path.join(data, fname)))
    check("run.json produced", os.path.exists(os.path.join(data, "run.json")))
    check("_DONE.flag produced", os.path.exists(os.path.join(data, "_DONE.flag")))
    manual = os.path.join(workdir, "MANUAL")
    for f in ("TeamsAdminExport.csv", "SPAdminSites.csv", "MDE_Devices.csv"):
        check(f"MANUAL {f} produced", os.path.exists(os.path.join(manual, f)))
    return data, manual


# ---------------------------------------------------------------------------
# Step 2: import into a copy of the workbook (real importer) + dummy prices
# ---------------------------------------------------------------------------
DUMMY_PRICES = {}


def set_dummy_prices(wb_path):
    """Type a deterministic dummy price on each Config pricing row via COM."""
    import win32com.client
    xl = win32com.client.DispatchEx("Excel.Application")
    xl.Visible = False
    xl.DisplayAlerts = False
    try:
        wb = xl.Workbooks.Open(wb_path)
        cfg = wb.Worksheets("Config")
        r = 3
        i = 0
        while True:
            name = cfg.Cells(r, 4).Value
            if name is None or str(name).strip() == "":
                break
            i += 1
            price = round(1.5 + (i % 7) * 2.25, 2)
            cfg.Cells(r, 5).Value = price
            DUMMY_PRICES[str(name)] = price
            r += 1
        wb.Save()
        wb.Close(False)
    finally:
        xl.Quit()


def step_import(workdir, data, manual, tag):
    print(f"\n== 2. Import into workbook ({tag}) ==")
    wb_path = os.path.join(workdir, f"wb_{tag}.xlsm")
    shutil.copy(XLSM, wb_path)
    if tag == "sample":
        set_dummy_prices(wb_path)
        check("dummy prices set", len(DUMMY_PRICES) > 10, str(len(DUMMY_PRICES)))
    ps(os.path.join(DIST, "Refresh-Workbook.ps1"),
       "-WorkbookPath", wb_path, "-DataPath", data, "-ManualPath", manual)
    return wb_path


# ---------------------------------------------------------------------------
# Step 3: KPI assertions with pandas
# ---------------------------------------------------------------------------
def step_assert(wb_path, data, manual):
    print("\n== 3. KPI assertions ==")
    import openpyxl
    wb = openpyxl.load_workbook(wb_path, data_only=True)

    errors = scan_errors(wb)
    check("zero formula errors (sample data)", len(errors) == 0, "; ".join(errors[:10]))

    ds = {}
    for key, (fname, _sheet, _max, cols) in contract.DATASETS.items():
        ds[key] = load_csv(os.path.join(data, fname), cols)

    with open(os.path.join(data, "run.json"), encoding="utf-8-sig") as fh:
        run_meta = json.load(fh)
    ref_date = dt.date.fromisoformat(run_meta["runStartUtc"][:10])
    va_d, a_d, la_d = 30, 90, 180

    u = ds["Users"]
    members = u[u.UserType == "Member"]
    check("Users total", num(cell(wb, "Dashboard!A6")) == len(members), f"{cell(wb,'Dashboard!A6')} vs {len(members)}")
    check("Users disabled", num(cell(wb, "Dashboard!D6")) == len(members[members.AccountEnabled == False]))  # noqa: E712
    check("Users licensed", num(cell(wb, "Dashboard!D7")) == len(members[members.IsLicensed == True]))  # noqa: E712
    check("Guest users", num(cell(wb, "Dashboard!D8")) == len(u[u.UserType == "Guest"]))
    synced = len(u[u.OnPremSynced == True])  # noqa: E712
    expected_idn = "Cloud-only" if synced == 0 else f"Hybrid ({synced} synced)"
    check("Identity type", cell(wb, "Dashboard!C10") == expected_idn, str(cell(wb, "Dashboard!C10")))

    g = ds["Groups"]
    check("Groups total", num(cell(wb, "Dashboard!A14")) == len(g))
    check("Groups security", num(cell(wb, "Dashboard!D14")) == len(g[g.Category == "Security"]))
    check("Groups M365", num(cell(wb, "Dashboard!D15")) == len(g[g.Category == "Microsoft 365"]))

    de = ds["DevicesEntra"]
    di = ds["DevicesIntune"]
    check("Devices total", num(cell(wb, "Dashboard!A19")) == len(de))
    seen_year = len(de[de.LastSignIn.notna() & (de.LastSignIn >= dt.date(ref_date.year, 1, 1))]) if len(de) else 0
    check("Devices seen this year", num(cell(wb, "Dashboard!A21")) == seen_year)
    stale = len(de[de.LastSignIn.notna() & (de.LastSignIn < ref_date - dt.timedelta(days=182))]) + len(de[de.LastSignIn.isna()])
    check("Devices stale", num(cell(wb, "Dashboard!B21")) == stale)
    check("Devices noncompliant", num(cell(wb, "Dashboard!C21")) == len(di[di.ComplianceState == "noncompliant"]))
    check("Devices not managed", num(cell(wb, "Dashboard!D21")) == len(de[de.IsManaged == False]))  # noqa: E712

    apps = ds["EnterpriseApps"]
    third = apps[apps.IsMicrosoftFirstParty == False]  # noqa: E712
    check("Enterprise apps", num(cell(wb, "Dashboard!A27")) == len(third))
    newest = third.sort_values("CreatedDateTime", ascending=False).DisplayName.iloc[0] if len(third) else ""
    check("Top app #1 is newest third-party", cell(wb, "Dashboard!A31") == newest,
          f"{cell(wb,'Dashboard!A31')} vs {newest}")

    mb = ds["Mailboxes"]
    check("Mailboxes total", num(cell(wb, "Dashboard!F6")) == len(mb))
    check("Total size", cell(wb, "Dashboard!H6") == f"{int(xlround(mb.SizeGB.sum()))} GB",
          f"{cell(wb,'Dashboard!H6')} vs {mb.SizeGB.sum()}")
    tcount = mb.RecipientTypeDetails.value_counts()

    def tc(name):
        return int(tcount.get(name, 0))
    check("Room mbx", num(cell(wb, "Dashboard!G8")) == tc("RoomMailbox") + tc("EquipmentMailbox"))
    check("Shared mbx", num(cell(wb, "Dashboard!G9")) == tc("SharedMailbox"))
    check("User mbx", num(cell(wb, "Dashboard!G10")) == tc("UserMailbox"))
    check("Group mbx", num(cell(wb, "Dashboard!G11")) == tc("GroupMailbox") + tc("TeamMailbox"))
    check("Booking mbx", num(cell(wb, "Dashboard!G12")) == tc("SchedulingMailbox"))
    breakdown = (num(cell(wb, "Dashboard!G8")) + num(cell(wb, "Dashboard!G9")) + num(cell(wb, "Dashboard!G10"))
                 + num(cell(wb, "Dashboard!G11")) + num(cell(wb, "Dashboard!G12")))
    check("Mailbox type breakdown sums to total", breakdown == len(mb))

    um = mb[mb.RecipientTypeDetails == "UserMailbox"]
    act = um[um.LastUserActionTime.notna()]
    va = len(act[act.LastUserActionTime >= ref_date - dt.timedelta(days=va_d)])
    a = len(act[(act.LastUserActionTime >= ref_date - dt.timedelta(days=a_d))
                & (act.LastUserActionTime < ref_date - dt.timedelta(days=va_d))])
    la = len(act[(act.LastUserActionTime >= ref_date - dt.timedelta(days=la_d))
                 & (act.LastUserActionTime < ref_date - dt.timedelta(days=a_d))])
    inact = len(um) - va - a - la
    check("Band VeryActive", num(cell(wb, "Dashboard!H15")) == va)
    check("Band Active", num(cell(wb, "Dashboard!H16")) == a)
    check("Band Low", num(cell(wb, "Dashboard!H17")) == la)
    check("Band Inactive", num(cell(wb, "Dashboard!H18")) == inact)
    check("Bands sum to UserMailbox count",
          num(cell(wb, "Dashboard!H15")) + num(cell(wb, "Dashboard!H16"))
          + num(cell(wb, "Dashboard!H17")) + num(cell(wb, "Dashboard!H18")) == len(um))
    if len(um):
        health = f"{xlround((va + a) / len(um) * 100, 1)}%".replace(".0%", "%")
        got = str(cell(wb, "Dashboard!H19")).replace(",", ".").replace(".0%", "%")
        check("Health score", got == health, f"{got} vs {health}")

    ea = ds["EmailActivity"]
    check("Traffic send", num(cell(wb, "Dashboard!F30")) == ea.Send.sum())
    check("Traffic receive", num(cell(wb, "Dashboard!G30")) == ea.Receive.sum())
    check("Traffic read", num(cell(wb, "Dashboard!H30")) == ea.Read.sum())

    check("DL count", num(cell(wb, "Dashboard!H32")) == len(ds["DistributionGroups"]))
    arch = len(mb[mb.ArchiveEnabled == True])  # noqa: E712
    expected_arch = f"Yes ({arch})" if arch > 0 else "No"
    check("Archives", cell(wb, "Dashboard!H33") == expected_arch, str(cell(wb, "Dashboard!H33")))
    check("Domains", num(cell(wb, "Dashboard!H34")) == len(ds["Domains"]))
    fwd = len(mb[mb.ForwardsExternally == True])  # noqa: E712
    check("External forwarding", num(cell(wb, "Dashboard!H35")) == fwd)

    sp = ds["SPOSites"]
    standalone = sp[sp.IsTeamsConnected == False]  # noqa: E712
    teams_sites = sp[sp.IsTeamsConnected == True]  # noqa: E712
    check("Site collections", num(cell(wb, "Dashboard!K6")) == len(standalone))
    check("SP storage standalone", abs(num(cell(wb, "Dashboard!M8")) - xlround(standalone.StorageUsedGB.sum(), 1)) < 0.11)
    sp_inact = len(standalone[standalone.LastActivityDate.isna()]) + \
        len(standalone[standalone.LastActivityDate.notna()
                       & (standalone.LastActivityDate < ref_date - dt.timedelta(days=la_d))])
    check("SP inactive sites", num(cell(wb, "Dashboard!M9")) == sp_inact)
    check("Teams sites", num(cell(wb, "Dashboard!K14")) == len(teams_sites))
    check("Teams sites storage", num(cell(wb, "Dashboard!M16")) == xlround(teams_sites.StorageUsedGB.sum()))
    check("Files", num(cell(wb, "Dashboard!M25")) == sp.FileCount.sum())
    m19 = len(teams_sites[teams_sites.LastActivityDate.notna()
                          & (teams_sites.LastActivityDate >= ref_date - dt.timedelta(days=va_d))])
    check("SP band VeryActive", num(cell(wb, "Dashboard!M19")) == m19)
    check("SP bands sum", num(cell(wb, "Dashboard!M19")) + num(cell(wb, "Dashboard!M20"))
          + num(cell(wb, "Dashboard!M21")) + num(cell(wb, "Dashboard!M22")) == len(teams_sites))
    check("Tenant sharing label", cell(wb, "Dashboard!M27") == "Anyone with link", str(cell(wb, "Dashboard!M27")))

    # SP admin manual sheet: external sharing counts
    spadm = pd.read_csv(os.path.join(manual, "SPAdminSites.csv"))
    m10 = len(spadm[(spadm.Template != "Team site") & (spadm["External sharing"] == "On")])
    m17 = len(spadm[(spadm.Template == "Team site") & (spadm["External sharing"] == "On")])
    check("SP sharing external (standalone)", num(cell(wb, "Dashboard!M10")) == m10)
    check("SP sharing external (team sites)", num(cell(wb, "Dashboard!M17")) == m17)

    tm = ds["Teams"]
    check("Teams count", num(cell(wb, "Dashboard!P6")) == len(tm))
    check("Teams public", num(cell(wb, "Dashboard!S6")) == len(tm[tm.Visibility == "Public"]))
    check("Teams private", num(cell(wb, "Dashboard!S7")) == len(tm[tm.Visibility == "Private"]))
    check("Teams members", num(cell(wb, "Dashboard!S15")) == tm.MemberCount.sum())

    tadm = pd.read_csv(os.path.join(manual, "TeamsAdminExport.csv"))
    check("Channels total", num(cell(wb, "Dashboard!P10")) ==
          tadm["Standard Channels"].sum() + tadm["Private Channels"].sum() + tadm["Shared Channels"].sum())
    check("Channels standard", num(cell(wb, "Dashboard!S10")) == tadm["Standard Channels"].sum())

    od = ds["OneDrive"]
    check("OD storage", num(cell(wb, "Dashboard!R29")) == xlround(od.StorageUsedGB.sum()))
    check("OD accounts", num(cell(wb, "Dashboard!R30")) == len(od))

    mfa = ds["MFA"]
    check("MFA registered", num(cell(wb, "Dashboard!U6")) == len(mfa[mfa.IsMfaRegistered == True]))  # noqa: E712
    mfa_upns = set(mfa[mfa.IsMfaRegistered == True].UserPrincipalName)  # noqa: E712
    not_reg = members[(members.AccountEnabled == True) & (members.IsLicensed == True)  # noqa: E712
                      & ~members.UserPrincipalName.isin(mfa_upns)]
    check("MFA not registered (licensed members)", num(cell(wb, "Dashboard!U9")) == len(not_reg))

    rl = ds["Roles"]
    for coord, role in (("X16", "Global Administrator"), ("X17", "SharePoint Administrator"),
                        ("X18", "Teams Administrator"), ("X19", "User Administrator"),
                        ("X20", "Exchange Administrator")):
        check(f"Role {role}", num(cell(wb, f"Dashboard!{coord}")) == len(rl[rl.RoleName == role]))
    check("Total privileged", num(cell(wb, "Dashboard!X21")) == len(rl))

    li = ds["Licenses"]
    li_priced = li.copy()
    li_priced["price"] = li_priced.FriendlyName.map(lambda n: DUMMY_PRICES.get(n, 0.0))
    total_month = (li_priced.Assigned * li_priced.price).sum()
    check("License total month", abs(num(cell(wb, "Dashboard!W33")) - total_month) < 0.01,
          f"{cell(wb,'Dashboard!W33')} vs {total_month}")
    check("License total year", abs(num(cell(wb, "Dashboard!W34")) - total_month * 12) < 0.01)
    assigned = li[li.Assigned > 0].sort_values("Assigned", ascending=False)
    if len(assigned):
        check("License slot 1 = biggest assigned", cell(wb, "Dashboard!U25") == assigned.FriendlyName.iloc[0],
              f"{cell(wb,'Dashboard!U25')} vs {assigned.FriendlyName.iloc[0]}")

    ca = ds["SecuritySettings"]
    if len(ca):
        expected_ca = f"{int(ca.CAPoliciesTotal.iloc[0])} ({int(ca.CAPoliciesEnabled.iloc[0])} enabled)"
        check("Conditional Access line", str(cell(wb, "Dashboard!W12")) == expected_ca, str(cell(wb, "Dashboard!W12")))

    # findings
    never = len(um[um.LastEmailSent.isna()])
    h37 = cell(wb, "Dashboard!H37") or ""
    check("Finding never-sent", (never == 0 and h37 == "") or (str(never) in h37), f"'{h37}' vs {never}")
    priv_nomfa = len(mfa[(mfa.IsAdmin == True) & (mfa.IsMfaRegistered == False)])  # noqa: E712
    h40 = cell(wb, "Dashboard!H40") or ""
    check("Finding privileged w/o MFA", (priv_nomfa == 0 and h40 == "") or (str(priv_nomfa) in h40), f"'{h40}' vs {priv_nomfa}")

    # MDE line
    mde = pd.read_csv(os.path.join(manual, "MDE_Devices.csv"))
    discovered = len(mde[mde["Onboarding Status"] == "Can be onboarded"])
    a25 = cell(wb, "Dashboard!A25") or ""
    check("MDE discovered line", (discovered == 0 and a25 == "") or (str(discovered) in a25), f"'{a25}' vs {discovered}")

    # RunInfo stamped
    check("RunInfo tenant", cell(wb, "RunInfo!B3") == run_meta["tenantDisplayName"], str(cell(wb, "RunInfo!B3")))
    check("RunInfo run date", cell(wb, "RunInfo!B4") is not None)

    wb.close()


# ---------------------------------------------------------------------------
# Step 4: empty-tenant case
# ---------------------------------------------------------------------------
def step_empty(workdir):
    print("\n== 4. Empty-data case ==")
    data = os.path.join(workdir, "DATA_EMPTY")
    manual = os.path.join(workdir, "MANUAL_EMPTY")
    os.makedirs(data, exist_ok=True)
    os.makedirs(manual, exist_ok=True)
    for key, (fname, _s, _m, cols) in contract.DATASETS.items():
        with open(os.path.join(data, fname), "w", encoding="utf-8-sig", newline="") as fh:
            fh.write(",".join(f'"{c}"' for c, _t in cols) + "\r\n")
    with open(os.path.join(data, "run.json"), "w", encoding="utf-8") as fh:
        json.dump({"tenantDisplayName": "Empty Tenant", "account": "t@e.local",
                   "tenantId": "0", "runStartUtc": "2026-07-08T00:00:00Z",
                   "concealedNames": False, "warningsCount": 0, "warnings": []}, fh)

    wb_path = step_import(workdir, data, manual, "empty")

    import openpyxl
    wb = openpyxl.load_workbook(wb_path, data_only=True)
    errors = scan_errors(wb)
    check("zero formula errors (empty data)", len(errors) == 0, "; ".join(errors[:10]))
    check("empty: users = 0", num(cell(wb, "Dashboard!A6")) == 0)
    check("empty: health = dash", cell(wb, "Dashboard!H19") == "-")
    check("empty: identity = dash", cell(wb, "Dashboard!C10") == "-")
    check("empty: channels = manual", cell(wb, "Dashboard!P10") == "manual")
    check("empty: findings blank", (cell(wb, "Dashboard!H37") or "") == "" and (cell(wb, "Dashboard!H38") or "") == "")
    wb.close()
    return wb_path


# ---------------------------------------------------------------------------
# Step 5: one-page PDF export
# ---------------------------------------------------------------------------
def step_pdf(wb_path, workdir):
    print("\n== 5. PDF export (one page) ==")
    import win32com.client
    pdf = os.path.join(workdir, "dashboard.pdf")
    xl = win32com.client.DispatchEx("Excel.Application")
    xl.Visible = False
    xl.DisplayAlerts = False
    try:
        wb = xl.Workbooks.Open(wb_path)
        dash = wb.Worksheets("Dashboard")
        dash.PageSetup.Zoom = False
        dash.PageSetup.FitToPagesWide = 1
        dash.PageSetup.FitToPagesTall = 1
        dash.ExportAsFixedFormat(0, pdf)
        wb.Close(False)
    finally:
        xl.Quit()
    with open(pdf, "rb") as fh:
        blob = fh.read()
    pages = len(re.findall(rb"/Type\s*/Page[^s]", blob))
    check("PDF has exactly 1 page", pages == 1, f"pages={pages}")


# ---------------------------------------------------------------------------
# Step 6: static consistency checks
# ---------------------------------------------------------------------------
def step_static():
    print("\n== 6. Static consistency ==")
    # charts bound to D_ sheets
    with zipfile.ZipFile(XLSM) as z:
        c1 = z.read("xl/charts/chart1.xml").decode("utf-8")
        c2 = z.read("xl/charts/chart2.xml").decode("utf-8")
    check("chart1 bound to D_EmailActivity",
          "D_EmailActivity!$B$2:$B$32" in c1 and "D_EmailActivity!$A$2:$A$32" in c1
          and "Email activity" not in c1)
    check("chart2 bound to D_TeamsUserActivity",
          "D_TeamsUserActivity!$D$2:$D$31" in c2 and "TeamsUserActivityCounts" not in c2)

    # dataset contract identical across the three implementations
    def parse_table(text):
        entries = re.findall(r'"(\w+\.csv\|D_\w+\|\d+\|[TDNB]+)"', text)
        entries += re.findall(r"'(\w+\.csv\|D_\w+\|\d+\|[TDNB]+)'", text)
        return sorted(set(entries))

    with open(os.path.join(DIST, "vba", "modImport.bas"), encoding="utf-8") as fh:
        vba_entries = parse_table(fh.read())
    with open(os.path.join(DIST, "Refresh-Workbook.ps1"), encoding="utf-8") as fh:
        ps_entries = parse_table(fh.read())
    py_entries = sorted(
        "{0}|{1}|{2}|{3}".format(f, s, m, "".join(t for _c, t in cols))
        for f, s, m, cols in contract.DATASETS.values())
    check("contract: VBA == contract.py", vba_entries == py_entries,
          f"vba={len(vba_entries)} py={len(py_entries)} diff={set(vba_entries) ^ set(py_entries)}")
    check("contract: Refresh-Workbook == contract.py", ps_entries == py_entries,
          f"ps={len(ps_entries)} diff={set(ps_entries) ^ set(py_entries)}")

    # SKU map identical between collector and workbook builder
    with open(os.path.join(DIST, "Collect-M365Audit.ps1"), encoding="utf-8") as fh:
        ps1 = fh.read()
    block = re.search(r"\$script:SkuFriendlyNames = @\{(.*?)\n\}", ps1, re.S).group(1)
    ps_map = dict(re.findall(r"'([^']+)'\s*=\s*'([^']+)'", block))
    check("SKU map: collector == builder", ps_map == SKU_MAP,
          f"only-collector={set(ps_map) - set(SKU_MAP)} only-builder={set(SKU_MAP) - set(ps_map)}")


# ---------------------------------------------------------------------------
def main():
    if not os.path.exists(XLSM):
        print("dist\\M365_Audit_Dashboard.xlsm not found - run tools\\build_workbook.py "
              "then Build-Workbook.ps1 first.")
        sys.exit(2)
    workdir = os.path.join(tempfile.gettempdir(), "m365audit_validate")
    if os.path.exists(workdir):
        shutil.rmtree(workdir)
    os.makedirs(workdir)
    print("workdir:", workdir)

    data, manual = step_collect(workdir)
    wb_path = step_import(workdir, data, manual, "sample")
    step_assert(wb_path, data, manual)
    empty_wb = step_empty(workdir)
    step_pdf(wb_path, workdir)
    step_static()

    print("\n" + "=" * 60)
    if FAILURES:
        print(f"RESULT: {len(FAILURES)} FAILURE(S)")
        for f in FAILURES:
            print("  -", f)
        sys.exit(1)
    print("RESULT: ALL CHECKS PASSED")
    sys.exit(0)


if __name__ == "__main__":
    main()
